"""报表导出：后端生成真实 Excel / PDF 文件并落库。

与诊断导出的同构设计：文件经 `asset_storage` 落地，导出记录落库可审计，
下载需登录。区别在于报表的内容是结构化数据（指标/趋势/明细），
PDF 直接用表格渲染，Excel 输出四个 sheet。
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import desc, select
from sqlalchemy.orm import Session

from app.db.models import ReportExportRecordModel
from app.schemas.reports import ReportExportRequest, ReportFormat, ReportOverview
from app.services.asset_storage import AssetStorageError, StoredAsset, put_bytes
from app.services.pdf_render import ensure_cid_font

SHANGHAI = ZoneInfo("Asia/Shanghai")

SHOP_LABELS = {
    "shop-all": "全部店铺",
    "shop-taobao-001": "柔肤旗舰店（淘宝）",
    "shop-douyin-001": "柔肤抖音专营店（抖店）",
    "shop-pdd-001": "柔肤品牌馆（拼多多）",
}

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F5EFF")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)

BOUNDARY_NOTES = [
    "导出文件由后端生成，与页面同源数据；指标口径与同步时间以系统内定义为准；",
    "当前数据源为演示种子数据，真实平台数据需完成店铺授权后同步；",
    "环比为与上一周期对比的演示值，历史快照接入后由快照计算；",
    "本报表为只读分析输出，不构成任何平台写操作指令。",
]


def _scope_label(shop_id: str) -> str:
    return SHOP_LABELS.get(shop_id, shop_id)


def _fmt_money(value: float) -> str:
    return f"¥{value:,.2f}"


def _autosize(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _style_header(sheet, row_index: int = 1) -> None:
    for cell in sheet[row_index]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


# ---------------------------------------------------------------- Excel


def render_report_excel(overview: ReportOverview, shop_id: str) -> bytes:
    workbook = Workbook()

    metrics_sheet = workbook.active
    metrics_sheet.title = "指标概览"
    scope_label = _scope_label(shop_id)
    template_name = next(
        (
            item.name
            for item in overview.templates
            if item.id == overview.selected_template
        ),
        str(overview.selected_template),
    )
    metrics_sheet.append([f"{template_name} · {scope_label}"])
    metrics_sheet["A1"].font = Font(bold=True, size=14)
    metrics_sheet.append(["报表模板", template_name])
    metrics_sheet.append(["数据范围", scope_label])
    metrics_sheet.append(
        ["数据更新时间", overview.data_updated_at.strftime("%Y-%m-%d %H:%M")]
    )
    metrics_sheet.append([])
    metrics_sheet.append(["指标", "数值", "环比(%)", "方向"])
    _style_header(metrics_sheet, metrics_sheet.max_row)
    for metric in overview.metrics:
        metrics_sheet.append(
            [metric.label, metric.display_value, metric.change, metric.direction]
        )
    _autosize(metrics_sheet, {"A": 18, "B": 16, "C": 14, "D": 12})

    trend_sheet = workbook.create_sheet("7日趋势")
    trend_sheet.append(["日期", "成交金额(元)", "支付订单", "广告消耗(元)", "退款金额(元)"])
    _style_header(trend_sheet)
    for point in overview.trend:
        trend_sheet.append(
            [point.date, point.revenue, point.orders, point.spend, point.refunds]
        )
    _autosize(trend_sheet, {"A": 12, "B": 16, "C": 12, "D": 16, "E": 16})

    rows_sheet = workbook.create_sheet("主体明细")
    rows_sheet.append(
        [
            "商品 / 主体",
            "店铺",
            "平台",
            "成交金额(元)",
            "支付订单",
            "广告消耗(元)",
            "ROI",
            "退款率(%)",
        ]
    )
    _style_header(rows_sheet)
    for row in overview.rows:
        rows_sheet.append(
            [
                row.name,
                row.shop_name,
                row.platform_label,
                row.revenue,
                row.orders,
                row.spend,
                row.roi,
                row.refund_rate,
            ]
        )
    _autosize(
        rows_sheet, {"A": 28, "B": 20, "C": 10, "D": 16, "E": 12, "F": 16, "G": 10, "H": 12}
    )

    boundary_sheet = workbook.create_sheet("口径与边界")
    boundary_sheet.append(["口径与边界声明"])
    boundary_sheet["A1"].font = Font(bold=True, size=12)
    for note in BOUNDARY_NOTES:
        boundary_sheet.append([note])
    _autosize(boundary_sheet, {"A": 90})

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- PDF


def render_report_pdf(overview: ReportOverview, shop_id: str) -> bytes:
    font = ensure_cid_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ReportBody", parent=styles["Normal"], fontName=font, fontSize=9.5, leading=14
    )
    title_style = ParagraphStyle(
        "ReportTitle", parent=body, fontSize=16, leading=20, spaceAfter=8
    )
    h2 = ParagraphStyle(
        "ReportH2", parent=body, fontSize=12, leading=16, spaceBefore=12, spaceAfter=6
    )
    scope_label = _scope_label(shop_id)
    template_name = next(
        (item.name for item in overview.templates if item.id == overview.selected_template),
        str(overview.selected_template),
    )

    def table(data: list[list[str]], col_width: float) -> Table:
        width = len(data[0])
        rendered = Table(
            [[Paragraph(str(cell), body) for cell in row] for row in data],
            colWidths=[col_width / width] * width,
            repeatRows=1,
        )
        rendered.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef4ff")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9d6ea")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        return rendered

    story: list = [
        Paragraph(f"{template_name} · {scope_label}", title_style),
        Paragraph(
            f"数据更新时间：{overview.data_updated_at.strftime('%Y-%m-%d %H:%M')}"
            f"　数据源：{'演示种子数据' if overview.is_demo else '实时授权店铺'}",
            body,
        ),
        Paragraph("一、指标概览", h2),
        table(
            [["指标", "数值", "环比", "方向"]]
            + [
                [
                    item.label,
                    item.display_value,
                    f"{item.change:+.1f}%",
                    "上升" if item.direction == "up" else "下降",
                ]
                for item in overview.metrics
            ],
            col_width=170 * mm,
        ),
        Paragraph("二、7 日趋势", h2),
        table(
            [["日期", "成交金额", "支付订单", "广告消耗", "退款金额"]]
            + [
                [
                    point.date,
                    _fmt_money(point.revenue),
                    str(point.orders),
                    _fmt_money(point.spend),
                    _fmt_money(point.refunds),
                ]
                for point in overview.trend
            ],
            col_width=170 * mm,
        ),
        Paragraph("三、主体明细", h2),
        table(
            [["商品 / 主体", "店铺", "平台", "成交金额", "ROI", "退款率"]]
            + [
                [
                    row.name,
                    row.shop_name,
                    row.platform_label,
                    _fmt_money(row.revenue),
                    f"{row.roi:.2f}",
                    f"{row.refund_rate:.1f}%",
                ]
                for row in overview.rows
            ],
            col_width=170 * mm,
        ),
        Paragraph("四、口径与边界声明", h2),
    ]
    for note in BOUNDARY_NOTES:
        story.append(Paragraph(f"• {note}", body))
    story.append(Spacer(1, 4 * mm))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"{template_name} · {scope_label}",
    )
    doc.build(story)
    return buffer.getvalue()


# ---------------------------------------------------------------- 任务


def _to_job(model: ReportExportRecordModel) -> dict:
    return {
        "id": model.id,
        "template_id": model.template_id,
        "template_name": model.template_name,
        "format": model.format,
        "status": model.status,
        "progress": 100 if model.status == "succeeded" else 0,
        "file_name": model.file_name,
        "file_size": model.file_size,
        "error_message": model.error_message,
        "created_at": model.created_at,
        "completed_at": model.completed_at,
        "is_demo": model.is_demo,
    }


def create_export(db: Session, organization_id: str, request: ReportExportRequest) -> dict:
    """生成真实报表文件并落库导出记录。"""
    from app.services.report_repository import get_report_overview

    overview = get_report_overview(
        db,
        shop_id=request.shop_id,
        template_id=request.template_id,
        metrics=",".join(request.metrics),
    )
    template_name = next(
        (item.name for item in overview.templates if item.id == request.template_id),
        str(request.template_id),
    )
    now = datetime.now(SHANGHAI)
    timestamp = now.strftime("%Y%m%d-%H%M")
    extension = "xlsx" if request.format == ReportFormat.EXCEL else "pdf"
    file_name = f"{template_name}_{_scope_label(request.shop_id)}_{timestamp}.{extension}"
    record_id = f"report-export-{uuid4().hex[:12]}"

    record = ReportExportRecordModel(
        id=record_id,
        organization_id=organization_id,
        template_id=request.template_id.value,
        template_name=template_name,
        format=request.format.value,
        shop_id=request.shop_id,
        status="succeeded",
        file_name=file_name,
        created_at=now,
        is_demo=True,
    )
    try:
        if request.format == ReportFormat.EXCEL:
            content = render_report_excel(overview, request.shop_id)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            content = render_report_pdf(overview, request.shop_id)
            content_type = "application/pdf"
        stored: StoredAsset = put_bytes(
            f"reports/{organization_id}/{record_id}.{extension}", content, content_type
        )
        record.object_key = stored.object_key
        record.file_size = stored.file_size
        record.completed_at = now
    except (AssetStorageError, ValueError) as exc:
        record.status = "failed"
        record.error_message = str(exc)[:500]
        record.completed_at = datetime.now(SHANGHAI)
        db.add(record)
        db.commit()
        raise
    db.add(record)
    db.commit()
    return _to_job(record)


def list_exports(db: Session, organization_id: str, limit: int = 20) -> list[dict]:
    rows = db.scalars(
        select(ReportExportRecordModel)
        .where(ReportExportRecordModel.organization_id == organization_id)
        .order_by(desc(ReportExportRecordModel.created_at))
        .limit(limit)
    ).all()
    return [_to_job(row) for row in rows]


def get_export(db: Session, organization_id: str, export_id: str) -> ReportExportRecordModel | None:
    record = db.get(ReportExportRecordModel, export_id)
    if record is None or record.organization_id != organization_id:
        return None
    return record
