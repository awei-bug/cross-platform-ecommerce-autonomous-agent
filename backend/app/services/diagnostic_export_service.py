"""诊断导出：后端生成真实 PDF / Excel 文件。

选型与边界：

- PDF 用 reportlab，中文走内置 CID 字体 `STSong-Light`，不依赖操作系统字体文件，
  本地与容器环境输出一致；
- Excel 用 openpyxl，问题清单 sheet 保留结构化字段，便于二次分析；
- 文件经 `asset_storage` 落地（local 或 MinIO），导出记录落库可审计，
  下载必须登录，对象键不含任何敏感信息。
"""
from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.db.models import DiagnosticExportModel, DiagnosticReportModel
from app.services.asset_storage import StoredAsset, put_bytes
from app.services.pdf_render import CID_FONT, ensure_cid_font

SHANGHAI = ZoneInfo("Asia/Shanghai")

_CID_FONT = CID_FONT

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F5EFF")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def _ensure_cid_font() -> None:
    ensure_cid_font()


def _md_inline(text: str) -> str:
    """把 Markdown 行内语法转成 reportlab 支持的标签。"""
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)


def _split_table_row(line: str) -> list[str]:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return [_md_inline(cell) for cell in cells]


def _markdown_to_story(markdown: str) -> list:
    """轻量 Markdown 渲染：标题、列表、表格、引用、段落。

    只覆盖模板报告与常见模型输出用到的结构，无法识别的行按普通段落处理，
    保证任何输入都能稳定出文件，而不是抛错。
    """
    _ensure_cid_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "CJKBody",
        parent=styles["Normal"],
        fontName=_CID_FONT,
        fontSize=9.5,
        leading=14,
    )
    h1 = ParagraphStyle("CJKH1", parent=body, fontSize=16, leading=20, spaceAfter=8)
    h2 = ParagraphStyle("CJKH2", parent=body, fontSize=13, leading=17, spaceBefore=10, spaceAfter=6)
    h3 = ParagraphStyle("CJKH3", parent=body, fontSize=11, leading=15, spaceBefore=8, spaceAfter=4)
    quote = ParagraphStyle(
        "CJKQuote",
        parent=body,
        textColor=colors.HexColor("#8a6d1f"),
        backColor=colors.HexColor("#fdf6e3"),
        borderPadding=4,
        spaceBefore=4,
        spaceAfter=4,
    )

    story: list = []
    table_buffer: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if not table_buffer:
            return
        width = max(len(row) for row in table_buffer)
        padded = [row + [""] * (width - len(row)) for row in table_buffer]
        table = Table(padded, colWidths=[170 * mm / width] * width, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), _CID_FONT),
                    ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef4ff")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9d6ea")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 4 * mm))
        table_buffer = []

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = _split_table_row(stripped)
            if all(re.fullmatch(r":?-{2,}:?", cell or "---") for cell in cells):
                continue  # 分隔行
            table_buffer.append(cells)
            continue
        flush_table()
        if not stripped:
            story.append(Spacer(1, 2 * mm))
            continue
        if stripped.startswith("### "):
            story.append(Paragraph(_md_inline(stripped[4:]), h3))
        elif stripped.startswith("## "):
            story.append(Paragraph(_md_inline(stripped[3:]), h2))
        elif stripped.startswith("# "):
            story.append(Paragraph(_md_inline(stripped[2:]), h1))
        elif stripped.startswith("> "):
            story.append(Paragraph(_md_inline(stripped[2:]), quote))
        elif stripped.startswith("- "):
            story.append(Paragraph("• " + _md_inline(stripped[2:]), body))
        elif re.match(r"^\d+\.\s", stripped):
            story.append(Paragraph(_md_inline(stripped), body))
        else:
            story.append(Paragraph(_md_inline(stripped), body))
    flush_table()
    return story


def render_report_pdf(report: dict) -> bytes:
    """把诊断报告渲染为 PDF 字节流。"""
    _ensure_cid_font()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=report["title"],
    )
    doc.build(_markdown_to_story(report["content"]))
    return buffer.getvalue()


def _autosize_columns(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def render_issues_workbook(overview: dict) -> bytes:
    """问题清单 Excel：健康分、问题明细、店铺对比三个 sheet。"""
    workbook = Workbook()

    health_sheet = workbook.active
    health_sheet.title = "健康分"
    health_sheet.append(["店铺健康分诊断"])
    health_sheet["A1"].font = Font(bold=True, size=13)
    health_sheet.append([])
    health_sheet.append(["范围", overview["shop_id"]])
    health_sheet.append(["数据源", overview["data_source"]])
    health_sheet.append(["算法版本", overview["health"]["algorithm_version"]])
    health_sheet.append(["综合健康分", overview["health"]["score"]])
    health_sheet.append([])
    health_sheet.append(["分项", "得分", "权重", "说明"])
    for cell in health_sheet[health_sheet.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for component in overview["health"]["components"]:
        health_sheet.append(
            [component["label"], component["score"], component["weight"], component["detail"]]
        )
    _autosize_columns(health_sheet, {"A": 16, "B": 10, "C": 10, "D": 60})

    issues_sheet = workbook.create_sheet("问题清单")
    issues_sheet.append(
        [
            "优先级",
            "严重度",
            "问题",
            "店铺",
            "商品",
            "类型",
            "置信度",
            "预计损失(元)",
            "损失口径",
            "紧迫度",
            "建议",
            "证据指标",
        ]
    )
    for cell in issues_sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for issue in overview["issues"]:
        evidence_text = "；".join(
            f"{item['metric']}={item['value']}({item['source']})" for item in issue["evidence"]
        )
        issues_sheet.append(
            [
                issue["priority"],
                issue["severity"],
                issue["title"],
                issue["shop_name"],
                issue["product_title"] or "",
                issue["kind"],
                issue["confidence"],
                issue["estimated_loss"],
                issue["loss_note"],
                issue["urgency"],
                issue["recommendation"],
                evidence_text,
            ]
        )
    _autosize_columns(
        issues_sheet,
        {
            "A": 10,
            "B": 10,
            "C": 34,
            "D": 18,
            "E": 22,
            "F": 14,
            "G": 10,
            "H": 14,
            "I": 40,
            "J": 10,
            "K": 44,
            "L": 50,
        },
    )

    compare_sheet = workbook.create_sheet("店铺对比")
    compare_sheet.append(
        ["店铺", "平台", "健康分", "SKU 数", "严重问题", "关注问题", "近 7 日销售额(元)", "数据源"]
    )
    for cell in compare_sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in overview["comparison"]:
        compare_sheet.append(
            [
                row["shop_name"],
                row["platform_label"],
                row["health_score"],
                row["product_count"],
                row["critical_issues"],
                row["warning_issues"],
                row["revenue_7d"],
                row["data_source"],
            ]
        )
    _autosize_columns(
        compare_sheet, {"A": 20, "B": 10, "C": 10, "D": 10, "E": 10, "F": 10, "G": 18, "H": 12}
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()




def create_export(
    db,
    organization_id: str,
    report_id: str,
    export_format: str,
) -> dict:
    """生成导出文件并落库导出记录。返回文件元信息。"""
    if export_format not in {"pdf", "xlsx"}:
        raise ValueError("导出格式仅支持 pdf 或 xlsx")
    report = db.get(DiagnosticReportModel, report_id)
    if report is None or report.organization_id != organization_id:
        raise LookupError("诊断报告不存在")

    now = datetime.now(SHANGHAI)
    safe_id = uuid4().hex[:12]
    if export_format == "pdf":
        content = render_report_pdf(
            {
                "title": report.title,
                "content": report.content,
                "snapshot": report.snapshot,
            }
        )
        content_type = "application/pdf"
        extension = "pdf"
    else:
        content = render_issues_workbook(
            {
                "shop_id": report.shop_id,
                "data_source": report.snapshot.get("data_source", "seeded"),
                "health": report.snapshot["health"],
                "issues": report.snapshot["issues"],
                "comparison": report.snapshot.get("comparison", []),
            }
        )
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        extension = "xlsx"

    object_key = f"diagnostics/{organization_id}/{report_id}/{safe_id}.{extension}"
    stored: StoredAsset = put_bytes(object_key, content, content_type)
    export = DiagnosticExportModel(
        id=f"diag-export-{safe_id}",
        report_id=report.id,
        organization_id=organization_id,
        format=export_format,
        object_key=stored.object_key,
        file_name=f"{report.title}.{extension}",
        file_size=stored.file_size,
        status="succeeded",
        is_demo=report.is_demo,
        created_at=now,
    )
    db.add(export)
    db.commit()
    return {
        "id": export.id,
        "report_id": export.report_id,
        "format": export.format,
        "file_name": export.file_name,
        "file_size": export.file_size,
        "status": export.status,
        "is_demo": export.is_demo,
        "created_at": export.created_at,
    }


def get_export(db, organization_id: str, export_id: str) -> DiagnosticExportModel | None:
    export = db.get(DiagnosticExportModel, export_id)
    if export is None or export.organization_id != organization_id:
        return None
    return export
